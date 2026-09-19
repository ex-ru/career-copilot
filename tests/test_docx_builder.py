"""
tests/test_docx_builder.py - Unit tests for Markdown to DOCX compilation.
"""

from pathlib import Path
from docx import Document
from core.docx_builder import DocxBuilder


def test_render_markdown_to_docx(tmp_path: Path):
    md_content = """# Заголовок Резюме
**Senior DevOps Engineer** | Remote

---

## Ключевые компетенции
- Kubernetes & Docker
- Terraform & Ansible
- Python & Go

| Технология | Опыт |
| :--- | :--- |
| Linux | 10 лет |
| K8s | 5 лет |

> Рекомендация руководителя: отличный специалист.
"""
    output_docx = tmp_path / "test_output.docx"
    result_path = DocxBuilder.render_markdown_to_docx(md_content, output_docx)

    assert result_path.exists()
    assert result_path.stat().st_size > 1000

    # Validate generated docx with python-docx
    doc = Document(result_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    assert any("Заголовок Резюме" in p for p in paragraphs)
    assert any("Kubernetes & Docker" in p for p in paragraphs)
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 3
